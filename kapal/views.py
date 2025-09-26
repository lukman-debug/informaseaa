from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse, HttpResponse
from django.contrib import messages
from django.views.decorators.csrf import csrf_exempt
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile
from django.utils.html import escape
from django.db import models
import json
import os
import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import datetime
from .models import Kapal

def kapal_list(request):
    """Display list of kapal with document information"""
    kapal_objects = Kapal.objects.all()
    context = {
        'kapal_list': kapal_objects,
        'title': 'Daftar Kapal'
    }
    return render(request, 'kapal/kapal_list.html', context)

def kapal_create(request):
    """Create new kapal with document upload"""
    if request.method == 'POST':
        try:
            # Get form data
            nama_kapal = request.POST.get('nama_kapal')
            nomor_registrasi = request.POST.get('nomor_registrasi')
            jenis_kapal = request.POST.get('jenis_kapal')
            ukuran_kapal = request.POST.get('ukuran_kapal')
            tahun_pembuatan = int(request.POST.get('tahun_pembuatan'))
            pemilik = request.POST.get('pemilik')
            pelabuhan_asal = request.POST.get('pelabuhan_asal')
            
            # Create kapal object
            kapal = Kapal.objects.create(
                nama_kapal=nama_kapal,
                nomor_registrasi=nomor_registrasi,
                jenis_kapal=jenis_kapal,
                ukuran_kapal=ukuran_kapal,
                tahun_pembuatan=tahun_pembuatan,
                pemilik=pemilik,
                pelabuhan_asal=pelabuhan_asal,
                dokumen=[]
            )
            
            # Handle document uploads
            documents = request.FILES.getlist('documents')
            dokumen_list = []
            
            for doc in documents:
                # Save file
                file_path = f'kapal_documents/{kapal.id}/{doc.name}'
                saved_path = default_storage.save(file_path, ContentFile(doc.read()))
                
                # Add to document list
                dokumen_list.append({
                    'name': doc.name,
                    'path': saved_path,
                    'size': doc.size,
                    'uploaded_at': datetime.now().isoformat()
                })
            
            # Update kapal with documents
            kapal.dokumen = dokumen_list
            kapal.save()
            
            messages.success(request, f'Kapal {nama_kapal} berhasil ditambahkan!')
            return redirect('kapal_list')
            
        except Exception as e:
            messages.error(request, f'Error: {str(e)}')
    
    return render(request, 'kapal/kapal_form.html', {'title': 'Tambah Kapal Baru'})

def kapal_detail(request, pk):
    """Display kapal details with documents"""
    kapal = get_object_or_404(Kapal, pk=pk)
    context = {
        'kapal': kapal,
        'title': f'Detail Kapal - {kapal.nama_kapal}'
    }
    return render(request, 'kapal/kapal_detail.html', context)

def kapal_edit(request, pk):
    """Edit kapal with document management"""
    kapal = get_object_or_404(Kapal, pk=pk)
    
    if request.method == 'POST':
        try:
            # Update kapal fields
            kapal.nama_kapal = request.POST.get('nama_kapal')
            kapal.nomor_registrasi = request.POST.get('nomor_registrasi')
            kapal.jenis_kapal = request.POST.get('jenis_kapal')
            kapal.ukuran_kapal = request.POST.get('ukuran_kapal')
            kapal.tahun_pembuatan = int(request.POST.get('tahun_pembuatan'))
            kapal.pemilik = request.POST.get('pemilik')
            kapal.pelabuhan_asal = request.POST.get('pelabuhan_asal')
            
            # Handle new document uploads
            new_documents = request.FILES.getlist('documents')
            dokumen_list = list(kapal.dokumen) if kapal.dokumen else []
            
            for doc in new_documents:
                # Save file
                file_path = f'kapal_documents/{kapal.id}/{doc.name}'
                saved_path = default_storage.save(file_path, ContentFile(doc.read()))
                
                # Add to document list
                dokumen_list.append({
                    'name': doc.name,
                    'path': saved_path,
                    'size': doc.size,
                    'uploaded_at': datetime.now().isoformat()
                })
            
            kapal.dokumen = dokumen_list
            kapal.save()
            
            messages.success(request, f'Kapal {kapal.nama_kapal} berhasil diperbarui!')
            return redirect('kapal_detail', pk=kapal.pk)
            
        except Exception as e:
            messages.error(request, f'Error: {str(e)}')
    
    context = {
        'kapal': kapal,
        'title': f'Edit Kapal - {kapal.nama_kapal}'
    }
    return render(request, 'kapal/kapal_form.html', context)

def kapal_delete(request, pk):
    """Delete kapal and its documents"""
    kapal = get_object_or_404(Kapal, pk=pk)
    
    if request.method == 'POST':
        try:
            nama_kapal = kapal.nama_kapal
            
            # Delete associated files
            for doc in kapal.dokumen:
                if default_storage.exists(doc['path']):
                    default_storage.delete(doc['path'])
            
            kapal.delete()
            messages.success(request, f'Kapal {nama_kapal} berhasil dihapus!')
            return redirect('kapal_list')
            
        except Exception as e:
            messages.error(request, f'Error: {str(e)}')
    
    context = {
        'kapal': kapal,
        'title': f'Hapus Kapal - {kapal.nama_kapal}'
    }
    return render(request, 'kapal/kapal_confirm_delete.html', context)

def kapal_api_list(request):
    """API endpoint for kapal data (for AJAX/search)"""
    search = request.GET.get('search', '')
    kapal_objects = Kapal.objects.all()
    
    if search:
        kapal_objects = kapal_objects.filter(
            models.Q(nama_kapal__icontains=search) |
            models.Q(nomor_registrasi__icontains=search) |
            models.Q(jenis_kapal__icontains=search) |
            models.Q(pemilik__icontains=search) |
            models.Q(pelabuhan_asal__icontains=search)
        )
    
    data = []
    for kapal in kapal_objects:
        document_names = [doc.get('name', '') for doc in kapal.dokumen] if kapal.dokumen else []
        data.append({
            'id': kapal.id,
            'nama_kapal': kapal.nama_kapal,
            'nomor_registrasi': kapal.nomor_registrasi,
            'jenis_kapal': kapal.jenis_kapal,
            'ukuran_kapal': kapal.ukuran_kapal,
            'tahun_pembuatan': kapal.tahun_pembuatan,
            'pemilik': kapal.pemilik,
            'pelabuhan_asal': kapal.pelabuhan_asal,
            'dokumen_count': kapal.get_dokumen_count(),
            'dokumen_names': document_names,
            'created_at': kapal.created_at.isoformat(),
        })
    
    return JsonResponse({'data': data})

@csrf_exempt
def upload_document(request):
    """Handle document upload via AJAX"""
    if request.method == 'POST':
        try:
            kapal_id = request.POST.get('kapal_id')
            kapal = get_object_or_404(Kapal, pk=kapal_id)
            
            files = request.FILES.getlist('files')
            uploaded_docs = []
            
            for file in files:
                # Validate file
                if file.size > 10 * 1024 * 1024:  # 10MB limit
                    return JsonResponse({'error': f'File {file.name} terlalu besar (max 10MB)'})
                
                # Save file
                file_path = f'kapal_documents/{kapal.id}/{file.name}'
                saved_path = default_storage.save(file_path, ContentFile(file.read()))
                
                doc_info = {
                    'name': file.name,
                    'path': saved_path,
                    'size': file.size,
                    'uploaded_at': datetime.now().isoformat()
                }
                
                uploaded_docs.append(doc_info)
            
            # Update kapal documents
            dokumen_list = list(kapal.dokumen) if kapal.dokumen else []
            dokumen_list.extend(uploaded_docs)
            kapal.dokumen = dokumen_list
            kapal.save()
            
            return JsonResponse({'success': True, 'documents': uploaded_docs})
            
        except Exception as e:
            return JsonResponse({'error': str(e)})
    
    return JsonResponse({'error': 'Invalid request method'})

def export_excel(request):
    """Export kapal data to Excel with document information"""
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Kapal"
    
    # Header style
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal='center')
    
    # Headers
    headers = [
        'No', 'Nama Kapal', 'Nomor Registrasi', 'Jenis Kapal', 
        'Ukuran Kapal', 'Tahun Pembuatan', 'Pemilik', 'Pelabuhan Asal',
        'Jumlah Dokumen', 'Nama Dokumen', 'Tanggal Dibuat'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
    
    # Data
    kapal_objects = Kapal.objects.all()
    row = 2
    
    for idx, kapal in enumerate(kapal_objects, 1):
        doc_names = ', '.join([doc.get('name', '') for doc in kapal.dokumen]) if kapal.dokumen else 'Tidak ada dokumen'
        
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=kapal.nama_kapal)
        ws.cell(row=row, column=3, value=kapal.nomor_registrasi)
        ws.cell(row=row, column=4, value=kapal.jenis_kapal)
        ws.cell(row=row, column=5, value=kapal.ukuran_kapal)
        ws.cell(row=row, column=6, value=kapal.tahun_pembuatan)
        ws.cell(row=row, column=7, value=kapal.pemilik)
        ws.cell(row=row, column=8, value=kapal.pelabuhan_asal)
        ws.cell(row=row, column=9, value=kapal.get_dokumen_count())
        ws.cell(row=row, column=10, value=doc_names)
        ws.cell(row=row, column=11, value=kapal.created_at.strftime('%Y-%m-%d %H:%M'))
        
        row += 1
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Create response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="data_kapal_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    
    wb.save(response)
    return response
